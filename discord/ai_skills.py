"""Bundled, progressively loaded skills and request-local resources.

This module has no bot/database imports so parsers and resource limits can be
tested without starting Discord or reading application configuration.
"""

from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path
from uuid import uuid4


MAX_ATTACHMENT_BYTES = 8_000_000
MAX_RESOURCE_BYTES = 16_000_000
MAX_RESOURCE_COUNT = 32
MAX_EXTRACTED_CHARS = 500_000
MAX_RESOURCE_SLICE = 16_000
SKILL_ROOT = Path(__file__).with_name("ai_skills")
SECRET_FILENAME = re.compile(r"(^\.env(?:\.|$)|\.(?:pem|key|p12|pfx)$|^(?:credentials|secrets?)(?:\.|$))", re.I)


@dataclass(frozen=True)
class Skill:
    name: str
    description: str
    tools: frozenset[str]
    instructions: str
    root: Path


class SkillRegistry:
    def __init__(self, root: Path = SKILL_ROOT):
        self.root = root.resolve()
        self.skills: dict[str, Skill] = {}
        for path in sorted(self.root.glob("*/SKILL.md")):
            if path.is_symlink() or path.parent.is_symlink():
                raise ValueError("Skill symlinks are not supported")
            raw = path.read_text(encoding="utf-8")
            if not raw.startswith("---\n"):
                raise ValueError(f"Missing skill metadata: {path.parent.name}")
            header, separator, body = raw[4:].partition("\n---\n")
            if not separator:
                raise ValueError("Unterminated skill metadata")
            metadata = {}
            for line in header.splitlines():
                key, colon, value = line.partition(":")
                if colon:
                    value = value.strip()
                    metadata[key.strip()] = json.loads(value) if value.startswith('"') else value
            name = metadata.get("name", "")
            description = metadata.get("description", "")
            if not re.fullmatch(r"[a-z0-9]+(?:-[a-z0-9]+)*", name) or name != path.parent.name or len(name) > 64:
                raise ValueError("Invalid skill name")
            if not description or len(description) > 1024 or len(raw) > 24_000:
                raise ValueError("Invalid skill description or oversized instructions")
            tools = frozenset(metadata.get("allowed-tools", "").split())
            self.skills[name] = Skill(name, description, tools, body.strip(), path.parent.resolve())

    @property
    def gated_tools(self) -> frozenset[str]:
        return frozenset(tool for skill in self.skills.values() for tool in skill.tools)

    def get(self, name: str) -> Skill:
        try:
            return self.skills[name]
        except KeyError:
            raise ValueError("Unknown skill; call list_skills for available names") from None

    def read_reference(self, name: str, relative_path: str) -> str:
        skill = self.get(name)
        relative = Path(relative_path)
        if SECRET_FILENAME.search(relative.name):
            raise ValueError("Secret or credential references are not supported")
        if relative.is_absolute() or ".." in relative.parts or relative.suffix.lower() != ".md":
            raise ValueError("Only Markdown references inside the loaded skill can be read")
        path = (skill.root / relative).resolve()
        if not path.is_relative_to(skill.root) or path.is_symlink():
            raise ValueError("Reference must stay inside its skill directory")
        if path.stat().st_size > 24_000:
            raise ValueError("Skill reference is too large")
        return path.read_text(encoding="utf-8")


@lru_cache(maxsize=1)
def get_skill_registry() -> SkillRegistry:
    return SkillRegistry()


@dataclass
class Resource:
    resource_id: str
    filename: str
    text: str
    media_type: str
    source: str

    def describe(self) -> dict:
        return {
            "resource_id": self.resource_id,
            "filename": self.filename,
            "media_type": self.media_type,
            "chars": len(self.text),
            "source": self.source,
        }


class ResourceStore:
    def __init__(self):
        self.items: dict[str, Resource] = {}
        self.total_bytes = 0

    def add(self, text: str, *, filename="result.txt", media_type="text/plain", source="tool") -> Resource:
        size = len(text.encode("utf-8"))
        if len(self.items) >= MAX_RESOURCE_COUNT or size + self.total_bytes > MAX_RESOURCE_BYTES:
            raise ValueError("Request resource limit reached; reuse existing resources")
        filename = re.split(r"[/\\]", str(filename))[-1][:120] or "result.txt"
        resource = Resource(f"res_{uuid4().hex[:16]}", filename, text, media_type, str(source)[:300])
        self.items[resource.resource_id] = resource
        self.total_bytes += size
        return resource

    def get(self, resource_id: str) -> Resource:
        try:
            return self.items[resource_id]
        except KeyError:
            raise ValueError("Resource is unavailable in this request") from None

    def read(self, resource_id: str, *, offset=0, limit=8000) -> dict:
        resource = self.get(resource_id)
        offset = max(0, int(offset))
        limit = max(1, min(MAX_RESOURCE_SLICE, int(limit)))
        text = resource.text[offset:offset + limit]
        return {
            **resource.describe(), "offset": offset, "text": text,
            "next_offset": offset + len(text) if offset + len(text) < len(resource.text) else None,
            "untrusted": True,
        }


@dataclass
class SkillSession:
    loaded: set[str] = field(default_factory=set)
    resources: ResourceStore = field(default_factory=ResourceStore)
    node: object | None = None
    node_slot: bool = False
    tool_calls: int = 0


def skill_session(context: dict) -> SkillSession:
    if "_skill_session" not in context:
        context["_skill_session"] = SkillSession()
    return context["_skill_session"]


def _check_archive(data: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > 2000 or sum(item.file_size for item in entries) > 32_000_000:
            raise ValueError("Document archive exceeds extraction limits")


def extract_attachment(data: bytes, filename: str) -> tuple[str, str, bool]:
    """Extract inert text only; never evaluate spreadsheet formulas or macros."""
    name = re.split(r"[/\\]", filename)[-1]
    if SECRET_FILENAME.search(name):
        raise ValueError("Secret or credential files are not supported")
    if len(data) > MAX_ATTACHMENT_BYTES:
        raise ValueError("Attachment exceeds the 8 MB limit")
    suffix = Path(name).suffix.lower()
    truncated = False
    if suffix in {".xlsx", ".docx"}:
        _check_archive(data)
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data), strict=False)
        if reader.is_encrypted:
            raise ValueError("Encrypted PDFs are not supported")
        parts = []
        remaining = MAX_EXTRACTED_CHARS
        for index, page in enumerate(reader.pages):
            if index >= 80 or remaining <= 0:
                truncated = True
                break
            text = page.extract_text() or ""
            parts.append(f"[Page {index + 1}]\n{text[:remaining]}")
            remaining -= len(text)
            truncated |= remaining < 0
        text = "\n\n".join(parts)
    elif suffix == ".docx":
        from xml.etree import ElementTree
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            root = ElementTree.fromstring(archive.read("word/document.xml"))
        ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        text = "\n".join("".join(node.itertext()) for node in root.iter(f"{ns}p"))
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        output = io.StringIO()
        writer = csv.writer(output)
        try:
            for sheet_index, sheet in enumerate(book):
                if sheet_index >= 10:
                    truncated = True
                    break
                writer.writerow([f"[Sheet: {sheet.title}]"])
                for index, row in enumerate(sheet.iter_rows(max_col=100, values_only=True)):
                    if index >= 5000 or output.tell() >= MAX_EXTRACTED_CHARS:
                        truncated = True
                        break
                    writer.writerow([str(value)[:4000] if value is not None else "" for value in row])
        finally:
            book.close()
        text = output.getvalue()
    else:
        allowed = {".txt", ".md", ".csv", ".tsv", ".json", ".jsonl", ".log", ".py", ".js", ".ts", ".html", ".xml", ".yaml", ".yml", ".sql"}
        if suffix not in allowed:
            raise ValueError("Supported files: text, CSV/TSV, JSON, PDF, DOCX, XLSX (no OCR)")
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = data.decode("utf-16") if data.startswith((b'\xff\xfe', b'\xfe\xff')) else data.decode("cp950")
            except UnicodeDecodeError:
                raise ValueError("Unsupported text encoding; use UTF-8") from None
        if "\x00" in text:
            raise ValueError("Binary attachments are not supported")
    truncated |= len(text) > MAX_EXTRACTED_CHARS
    return text[:MAX_EXTRACTED_CHARS], "text/plain", truncated
